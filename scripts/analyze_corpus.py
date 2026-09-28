#!/usr/bin/env python3
"""Generate a structural Markdown profile for the local document corpus."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import re
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path, PurePosixPath
from typing import TYPE_CHECKING, Any

from defusedxml import ElementTree


if TYPE_CHECKING:
    from collections.abc import Iterable


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".docx",
    ".md",
    ".markdown",
    ".pdf",
    ".pptx",
    ".txt",
    ".xlsx",
}
OFFICE_MARKERS = {
    ".docx": "word/document.xml",
    ".pptx": "ppt/presentation.xml",
    ".xlsx": "xl/workbook.xml",
}
MAX_OFFICE_EXPANDED_BYTES = 1024 * 1024 * 1024
MAX_OFFICE_EXPANSION_RATIO = 100
LOW_CONFIDENCE_THRESHOLD = 0.5
CONFIGURATION = {
    "docling": "2.129.0",
    "ocr_engine": "tesseract-cli",
    "ocr_languages": ["eng"],
    "force_full_page_ocr": False,
    "table_mode": "accurate",
    "table_cell_matching": True,
    "remote_services": False,
    "external_plugins": False,
    "picture_description": False,
    "chart_extraction": False,
    "code_enrichment": False,
    "formula_enrichment": False,
}


@dataclass
class UnitMetrics:
    number: int
    label: str
    words: int = 0
    tables: int = 0
    pictures: int = 0
    nonempty_cells: int = 0
    rows_used: int = 0
    columns_used: int = 0
    native_cells: int = 0
    ocr_cells: int = 0
    category: str = "not-applicable"
    visibility: str = "visible"


@dataclass
class FileRecord:
    path: Path
    relative_path: str
    extension: str
    size: int = 0
    sha256: str = ""
    integrity: str = "unchecked"
    status: str = "pending"
    duration_seconds: float = 0.0
    words: int = 0
    pages: int = 0
    tables: int = 0
    pictures: int = 0
    layout_items: int = 0
    comments: int = 0
    formulas: int = 0
    formulas_without_cached_values: int = 0
    confidence: dict[str, float] = field(default_factory=dict)
    units: list[UnitMetrics] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def supported(self) -> bool:
        return self.extension in SUPPORTED_EXTENSIONS


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data"))
    parser.add_argument(
        "--output", type=Path, default=Path("output/corpus-analysis.md")
    )
    parser.add_argument("--max-files", type=positive_integer)
    return parser.parse_args(argv)


def positive_integer(value: str) -> int:
    parsed = int(value)
    if parsed < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return parsed


def discover_files(root: Path) -> list[Path]:
    """Return regular, non-hidden files without following directory symlinks."""
    if not root.is_dir():
        raise ValueError(f"input directory does not exist: {root}")

    discovered: list[Path] = []
    for current, directories, filenames in os.walk(root, followlinks=False):
        current_path = Path(current)
        directories[:] = sorted(
            name
            for name in directories
            if not name.startswith(".")
            and name.casefold() != "certbot"
            and not (current_path / name).is_symlink()
        )
        for filename in sorted(filenames):
            path = current_path / filename
            if filename.startswith(".") or path.is_symlink() or not path.is_file():
                continue
            discovered.append(path)

    return sorted(discovered, key=lambda path: path.relative_to(root).as_posix().casefold())


def select_files(paths: list[Path], max_files: int | None) -> list[Path]:
    if max_files is None:
        return paths
    supported = [path for path in paths if path.suffix.casefold() in SUPPORTED_EXTENSIONS]
    return supported[:max_files]


def inspect_file(path: Path, root: Path) -> FileRecord:
    record = FileRecord(
        path=path,
        relative_path=path.relative_to(root).as_posix(),
        extension=path.suffix.casefold(),
        size=path.stat().st_size,
    )
    record.sha256 = sha256_file(path)

    if not record.supported:
        record.integrity = "not-applicable"
        record.status = "skipped"
        record.warnings.append("Unsupported file type")
        return record

    if record.size == 0:
        record.integrity = "failed"
        record.status = "failure"
        record.errors.append("Empty file")
        return record

    if record.extension == ".pdf":
        with path.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                record.integrity = "failed"
                record.status = "failure"
                record.errors.append("PDF signature does not match the extension")
                return record

    if record.extension in OFFICE_MARKERS:
        validate_office_container(record)
        if record.integrity == "failed":
            record.status = "failure"
            return record
        if record.extension == ".xlsx":
            (
                record.formulas,
                record.formulas_without_cached_values,
            ) = inspect_xlsx_formulas(path)

    record.integrity = "warning" if record.warnings else "valid"
    return record


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_office_container(record: FileRecord) -> None:
    try:
        with zipfile.ZipFile(record.path) as archive:
            members = archive.infolist()
            names = {member.filename for member in members}
            if OFFICE_MARKERS[record.extension] not in names:
                record.integrity = "failed"
                record.errors.append("Office container is missing its required document part")
                return

            if any(is_unsafe_archive_path(member.filename) for member in members):
                record.integrity = "failed"
                record.errors.append("Office container contains an unsafe archive path")
                return

            expanded = sum(member.file_size for member in members)
            compressed = sum(member.compress_size for member in members)
            ratio = expanded / max(compressed, 1)
            if expanded > MAX_OFFICE_EXPANDED_BYTES:
                record.warnings.append("Office container expands beyond 1 GiB")
            if ratio > MAX_OFFICE_EXPANSION_RATIO:
                record.warnings.append("Office container has an expansion ratio above 100:1")
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile):
        record.integrity = "failed"
        record.errors.append("Invalid Office ZIP container")


def is_unsafe_archive_path(name: str) -> bool:
    path = PurePosixPath(name)
    return path.is_absolute() or ".." in path.parts


def inspect_xlsx_formulas(path: Path) -> tuple[int, int]:
    """Count formulas and formulas lacking cached values without loading cells."""
    formulas = 0
    missing_cached_values = 0
    with zipfile.ZipFile(path) as archive:
        worksheet_names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
        for worksheet_name in worksheet_names:
            with archive.open(worksheet_name) as stream:
                for _, element in ElementTree.iterparse(stream, events=("end",)):
                    if not element.tag.endswith("}c"):
                        continue
                    formula = next(
                        (child for child in element if child.tag.endswith("}f")), None
                    )
                    if formula is not None:
                        formulas += 1
                        cached = next(
                            (child for child in element if child.tag.endswith("}v")),
                            None,
                        )
                        if cached is None or not (cached.text or "").strip():
                            missing_cached_values += 1
                    element.clear()
    return formulas, missing_cached_values


def build_converter() -> Any:
    from docling.datamodel.accelerator_options import (  # noqa: PLC0415
        AcceleratorDevice,
        AcceleratorOptions,
    )
    from docling.datamodel.base_models import InputFormat  # noqa: PLC0415
    from docling.datamodel.pipeline_options import (  # noqa: PLC0415
        PdfPipelineOptions,
        TableFormerMode,
        TableStructureOptions,
        TesseractCliOcrOptions,
    )
    from docling.document_converter import (  # noqa: PLC0415
        DocumentConverter,
        PdfFormatOption,
    )

    pipeline_options = PdfPipelineOptions(
        accelerator_options=AcceleratorOptions(device=AcceleratorDevice.CPU),
        enable_remote_services=False,
        allow_external_plugins=False,
        artifacts_path=os.environ.get("DOCLING_ARTIFACTS_PATH", "/opt/docling/models"),
        do_picture_classification=False,
        do_picture_description=False,
        do_chart_extraction=False,
        do_table_structure=True,
        table_structure_options=TableStructureOptions(
            mode=TableFormerMode.ACCURATE,
            do_cell_matching=True,
        ),
        do_ocr=True,
        ocr_options=TesseractCliOcrOptions(lang=["eng"]),
        do_code_enrichment=False,
        do_formula_enrichment=False,
    )
    allowed_formats = [
        InputFormat.PDF,
        InputFormat.DOCX,
        InputFormat.XLSX,
        InputFormat.CSV,
        InputFormat.PPTX,
        InputFormat.MD,
    ]
    return DocumentConverter(
        allowed_formats=allowed_formats,
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)},
    )


def analyze_record(record: FileRecord, converter: Any) -> None:
    started = time.monotonic()
    try:
        source: Any = record.path
        if record.extension == ".markdown":
            from io import BytesIO  # noqa: PLC0415

            from docling_core.types.io import DocumentStream  # noqa: PLC0415

            source = DocumentStream(
                name=f"{record.path.stem}.md",
                stream=BytesIO(record.path.read_bytes()),
            )
        result = converter.convert(source, raises_on_error=False)
        record.status = result.status.value
        record.errors.extend(format_docling_errors(result.errors))
        collect_document_metrics(record, result)
        collect_confidence(record, result.confidence)
        add_review_warnings(record)
    except Exception as error:  # Document failures must not abort the corpus run.
        record.status = "failure"
        record.errors.append(format_exception(error, record))
    finally:
        record.duration_seconds = time.monotonic() - started


def format_docling_errors(errors: Iterable[Any]) -> list[str]:
    formatted = []
    for error in errors:
        category = getattr(getattr(error, "category", None), "value", "unknown")
        module = getattr(error, "module_name", "unknown module")
        page = getattr(error, "page_no", None)
        location = f", page {page}" if page is not None else ""
        formatted.append(f"{category} in {module}{location}")
    return formatted


def format_exception(error: Exception, record: FileRecord) -> str:
    message = " ".join(str(error).split())
    message = message.replace(str(record.path), record.relative_path)
    if not message:
        message = "conversion failed"
    return f"{type(error).__name__}: {message[:300]}"


def collect_document_metrics(record: FileRecord, result: Any) -> None:
    document = result.document
    record.pages = len(document.pages)
    record.tables = len(document.tables)
    record.pictures = len(document.pictures)
    record.layout_items = len(document.texts) + record.tables + record.pictures
    record.comments = sum(
        len(getattr(item, "comments", []))
        for item in [*document.texts, *document.tables, *document.pictures]
    )

    if record.extension not in {".pdf", ".pptx", ".xlsx"}:
        record.words = document_word_count(document)
        return

    page_words = words_by_page(document)
    page_tables = item_counts_by_page(document.tables)
    page_pictures = item_counts_by_page(document.pictures)
    page_cells = table_cells_by_page(document.tables)
    page_extents = table_extents_by_page(document.tables)
    page_ocr = ocr_cells_by_page(result.pages)
    unit_names, visibility = unit_metadata(record, document)

    all_pages = sorted(
        set(document.pages)
        | set(page_words)
        | set(page_tables)
        | set(page_pictures)
        | set(page_cells)
    )
    record.units = []
    for page_number in all_pages:
        native_cells, ocr_cells = page_ocr.get(page_number, (0, 0))
        rows_used, columns_used = page_extents.get(page_number, (0, 0))
        record.units.append(
            UnitMetrics(
                number=page_number,
                label=unit_names.get(page_number, unit_label(record, page_number)),
                words=page_words.get(page_number, 0),
                tables=page_tables.get(page_number, 0),
                pictures=page_pictures.get(page_number, 0),
                nonempty_cells=page_cells.get(page_number, 0),
                rows_used=rows_used,
                columns_used=columns_used,
                native_cells=native_cells,
                ocr_cells=ocr_cells,
                category=(
                    classify_page(native_cells, ocr_cells)
                    if record.extension == ".pdf"
                    else "not-applicable"
                ),
                visibility=visibility.get(page_number, "visible"),
            )
        )

    if record.units:
        record.words = sum(unit.words for unit in record.units)
    else:
        record.words = document_word_count(document)


def words_by_page(document: Any) -> dict[int, int]:
    counts: defaultdict[int, int] = defaultdict(int)
    for item in document.texts:
        add_words_to_page(counts, item, getattr(item, "text", ""))
    for table in document.tables:
        text = " ".join(cell.text for cell in table.data.table_cells if cell.text)
        add_words_to_page(counts, table, text)
    return dict(counts)


def add_words_to_page(counts: defaultdict[int, int], item: Any, text: str) -> None:
    page_numbers = sorted({provenance.page_no for provenance in item.prov})
    if page_numbers:
        counts[page_numbers[0]] += word_count(text)


def item_counts_by_page(items: Iterable[Any]) -> dict[int, int]:
    counts: defaultdict[int, int] = defaultdict(int)
    for item in items:
        for page_number in {provenance.page_no for provenance in item.prov}:
            counts[page_number] += 1
    return dict(counts)


def table_cells_by_page(tables: Iterable[Any]) -> dict[int, int]:
    counts: defaultdict[int, int] = defaultdict(int)
    for table in tables:
        nonempty = sum(
            1 for cell in table.data.table_cells if (cell.text or "").strip()
        )
        page_numbers = sorted({provenance.page_no for provenance in table.prov})
        if page_numbers:
            counts[page_numbers[0]] += nonempty
    return dict(counts)


def table_extents_by_page(tables: Iterable[Any]) -> dict[int, tuple[int, int]]:
    extents: dict[int, tuple[int, int]] = {}
    for table in tables:
        for provenance in table.prov:
            rows, columns = extents.get(provenance.page_no, (0, 0))
            extents[provenance.page_no] = (
                max(rows, math.ceil(provenance.bbox.b)),
                max(columns, math.ceil(provenance.bbox.r)),
            )
    return extents


def document_word_count(document: Any) -> int:
    text_words = sum(word_count(item.text) for item in document.texts)
    table_words = sum(
        word_count(cell.text)
        for table in document.tables
        for cell in table.data.table_cells
    )
    return text_words + table_words


def word_count(text: str) -> int:
    return len(text.split())


def ocr_cells_by_page(pages: Iterable[Any]) -> dict[int, tuple[int, int]]:
    counts: dict[int, tuple[int, int]] = {}
    for page in pages:
        page_number = page.page_no
        seen: set[int] = set()
        native = 0
        ocr = 0
        assembled = getattr(page, "assembled", None)
        for element in getattr(assembled, "elements", []):
            cluster = getattr(element, "cluster", None)
            for cell in iter_cluster_cells(cluster):
                identifier = getattr(cell, "index", id(cell))
                if identifier in seen:
                    continue
                seen.add(identifier)
                if getattr(cell, "from_ocr", False):
                    ocr += 1
                else:
                    native += 1
        counts[page_number] = (native, ocr)
    return counts


def iter_cluster_cells(cluster: Any) -> Iterable[Any]:
    if cluster is None:
        return
    yield from getattr(cluster, "cells", [])
    for child in getattr(cluster, "children", []):
        yield from iter_cluster_cells(child)


def classify_page(native_cells: int, ocr_cells: int) -> str:
    if native_cells and ocr_cells:
        return "mixed"
    if ocr_cells:
        return "ocr-only scan candidate"
    if native_cells:
        return "native-only"
    return "empty"


def unit_label(record: FileRecord, number: int) -> str:
    if record.extension == ".xlsx":
        return f"Sheet {number}"
    if record.extension == ".pptx":
        return f"Slide {number}"
    return f"Page {number}"


def unit_metadata(
    record: FileRecord, document: Any
) -> tuple[dict[int, str], dict[int, str]]:
    if record.extension != ".xlsx":
        return {}, {}

    names = [
        group.name
        for group in document.groups
        if getattr(getattr(group, "label", None), "value", None) == "sheet"
    ]
    labels = dict(enumerate(names, 1))
    try:
        visibility_states = workbook_visibility(record.path)
    except Exception:
        visibility_states = {}
        record.warnings.append("Worksheet visibility metadata could not be read")
    visibility = {
        index: visibility_states.get(name, "unknown")
        for index, name in enumerate(names, 1)
    }
    return labels, visibility


def workbook_visibility(path: Path) -> dict[str, str]:
    from openpyxl import load_workbook  # noqa: PLC0415

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: sheet.sheet_state for sheet in workbook.worksheets}
    finally:
        workbook.close()


def collect_confidence(record: FileRecord, confidence: Any) -> None:
    for name in ("parse_score", "layout_score", "table_score", "ocr_score"):
        value = getattr(confidence, name, math.nan)
        if isinstance(value, (int, float)) and math.isfinite(value):
            record.confidence[name.removesuffix("_score")] = float(value)


def add_review_warnings(record: FileRecord) -> None:
    if record.status == "partial_success":
        record.warnings.append("Docling reported partial success")
    if record.words == 0:
        record.warnings.append("No words were extracted")
    if any(value < LOW_CONFIDENCE_THRESHOLD for value in record.confidence.values()):
        record.warnings.append("One or more confidence scores are below 0.50")
    if record.formulas:
        record.warnings.append(
            f"{record.formulas} formula cells use cached values and were not recalculated"
        )
    if record.formulas_without_cached_values:
        record.warnings.append(
            f"{record.formulas_without_cached_values} formula cells have no cached value"
        )
    hidden = [
        unit.label
        for unit in record.units
        if unit.visibility in {"hidden", "veryHidden"}
    ]
    if hidden:
        record.warnings.append(
            f"Hidden spreadsheet sheets require review: {', '.join(hidden)}"
        )


def duplicate_groups(records: list[FileRecord]) -> list[list[FileRecord]]:
    grouped: defaultdict[str, list[FileRecord]] = defaultdict(list)
    for record in records:
        grouped[record.sha256].append(record)
    return [group for group in grouped.values() if len(group) > 1]


def filename_collisions(records: list[FileRecord]) -> list[list[FileRecord]]:
    grouped: defaultdict[str, list[FileRecord]] = defaultdict(list)
    for record in records:
        grouped[Path(record.relative_path).name.casefold()].append(record)
    return [group for group in grouped.values() if len(group) > 1]


def render_report(
    records: list[FileRecord],
    input_label: str,
    max_files: int | None,
    started_at: datetime,
    elapsed_seconds: float,
) -> str:
    lines = [
        "# QueryMaster Corpus Analysis",
        "",
        "## Run configuration",
        "",
        "| Setting | Value |",
        "|---|---|",
        f"| Started | {escape_markdown(started_at.isoformat())} |",
        f"| Run type | {'Limited' if max_files else 'Full corpus'} |",
        f"| Maximum files | {max_files or 'None'} |",
        f"| Input | {escape_markdown(input_label)} |",
        f"| Docling | {escape_markdown(version('docling'))} |",
        f"| Tesseract | {escape_markdown(tesseract_version())} |",
        f"| Python | {escape_markdown(platform.python_version())} |",
        f"| Architecture | {escape_markdown(platform.machine())} |",
        "| Network during analysis | Disabled |",
        f"| Configuration fingerprint | `{configuration_fingerprint()}` |",
        f"| Model inventory fingerprint | `{model_inventory_fingerprint()}` |",
        f"| Elapsed seconds | {elapsed_seconds:.2f} |",
        "",
    ]
    lines.extend(render_inventory(records))
    lines.extend(render_integrity(records))
    lines.extend(render_identity(records))
    lines.extend(render_docling_summary(records))
    lines.extend(render_ocr_summary(records))
    lines.extend(render_review_queue(records))
    lines.extend(render_appendix(records))
    return "\n".join(lines).rstrip() + "\n"


def render_inventory(records: list[FileRecord]) -> list[str]:
    formats = Counter(record.extension.removeprefix(".").upper() or "NO EXTENSION" for record in records)
    lines = [
        "## Inventory summary",
        "",
        f"- Files: {len(records)}",
        f"- Total bytes: {sum(record.size for record in records):,}",
        "",
        "| Format | Files | Bytes |",
        "|---|---:|---:|",
    ]
    for extension, count in sorted(formats.items()):
        size = sum(
            record.size
            for record in records
            if (record.extension.removeprefix(".").upper() or "NO EXTENSION") == extension
        )
        lines.append(f"| {escape_markdown(extension)} | {count} | {size:,} |")
    return [*lines, ""]


def render_integrity(records: list[FileRecord]) -> list[str]:
    counts = Counter(record.integrity for record in records)
    lines = [
        "## Integrity findings",
        "",
        "| Result | Files |",
        "|---|---:|",
    ]
    for result, count in sorted(counts.items()):
        lines.append(f"| {escape_markdown(result)} | {count} |")
    return [*lines, ""]


def render_identity(records: list[FileRecord]) -> list[str]:
    lines = ["## Identity findings", "", "### Byte-identical files", ""]
    duplicates = duplicate_groups(records)
    if not duplicates:
        lines.append("None.")
    else:
        for group in duplicates:
            paths = ", ".join(f"`{escape_markdown(item.relative_path)}`" for item in group)
            lines.append(f"- `{group[0].sha256[:12]}…`: {paths}")
    lines.extend(["", "### Lowercase filename collisions", ""])
    collisions = filename_collisions(records)
    if not collisions:
        lines.append("None.")
    else:
        for group in collisions:
            paths = ", ".join(f"`{escape_markdown(item.relative_path)}`" for item in group)
            lines.append(f"- {paths}")
    return [*lines, ""]


def render_docling_summary(records: list[FileRecord]) -> list[str]:
    statuses = Counter(record.status for record in records)
    durations = [record.duration_seconds for record in records if record.duration_seconds]
    lines = [
        "## Docling summary",
        "",
        "| Status | Files |",
        "|---|---:|",
    ]
    for status, count in sorted(statuses.items()):
        lines.append(f"| {escape_markdown(status)} | {count} |")
    lines.extend(
        [
            "",
            f"- Extracted words: {sum(record.words for record in records):,}",
            f"- Pages/sheets/slides: {sum(record.pages for record in records):,}",
            f"- Tables: {sum(record.tables for record in records):,}",
            f"- Pictures: {sum(record.pictures for record in records):,}",
            f"- Layout items: {sum(record.layout_items for record in records):,}",
            f"- Median document duration: {statistics.median(durations):.2f} seconds"
            if durations
            else "- Median document duration: not available",
            "",
        ]
    )
    return lines


def render_ocr_summary(records: list[FileRecord]) -> list[str]:
    categories = Counter(
        unit.category
        for record in records
        if record.extension == ".pdf"
        for unit in record.units
    )
    lines = [
        "## PDF and OCR summary",
        "",
        "| Page classification | Pages |",
        "|---|---:|",
    ]
    lines.extend(
        f"| {category} | {categories[category]} |"
        for category in ("native-only", "ocr-only scan candidate", "mixed", "empty")
    )
    return [*lines, ""]


def render_review_queue(records: list[FileRecord]) -> list[str]:
    flagged = [
        record
        for record in records
        if record.errors or record.warnings or record.status in {"failure", "partial_success"}
    ]
    lines = [
        "## Review queue",
        "",
        "| File | Status | Reasons |",
        "|---|---|---|",
    ]
    if not flagged:
        lines.append("| — | — | No review findings |")
    for record in flagged:
        reasons = "; ".join([*record.errors, *record.warnings]) or "Review required"
        lines.append(
            f"| {escape_markdown(record.relative_path)} | "
            f"{escape_markdown(record.status)} | {escape_markdown(reasons)} |"
        )
    return [*lines, ""]


def render_appendix(records: list[FileRecord]) -> list[str]:
    lines = ["## Per-file appendix", ""]
    for record in records:
        lines.extend(
            [
                f"### {escape_markdown(record.relative_path)}",
                "",
                f"- Format: {escape_markdown(record.extension.removeprefix('.').upper())}",
                f"- Size: {record.size:,} bytes",
                f"- SHA-256: `{record.sha256[:12]}…`",
                f"- Integrity: {escape_markdown(record.integrity)}",
                f"- Conversion: {escape_markdown(record.status)}",
                f"- Duration: {record.duration_seconds:.2f} seconds",
                f"- Words: {record.words:,}",
                f"- Tables: {record.tables:,}",
                f"- Pictures: {record.pictures:,}",
                f"- Comments: {record.comments:,}",
            ]
        )
        if record.extension == ".xlsx":
            lines.extend(
                [
                    f"- Formula cells: {record.formulas:,}",
                    "- Formula cells without cached values: "
                    f"{record.formulas_without_cached_values:,}",
                ]
            )
        if record.confidence:
            scores = ", ".join(
                f"{name}={value:.3f}" for name, value in sorted(record.confidence.items())
            )
            lines.append(f"- Confidence: {scores}")
        if record.warnings:
            lines.append(f"- Warnings: {escape_markdown('; '.join(record.warnings))}")
        if record.errors:
            lines.append(f"- Errors: {escape_markdown('; '.join(record.errors))}")
        lines.append("")
        if record.units:
            label = unit_column_label(record.extension)
            lines.extend(
                [
                    f"| {label} | Visibility | Words | Tables | Pictures | Rows used | Columns used | Non-empty cells | OCR classification |",
                    "|---|---|---:|---:|---:|---:|---:|---:|---|",
                ]
            )
            lines.extend(
                (
                    f"| {escape_markdown(unit.label)} | {escape_markdown(unit.visibility)} | "
                    f"{unit.words} | {unit.tables} | {unit.pictures} | "
                    f"{unit.rows_used} | {unit.columns_used} | {unit.nonempty_cells} | "
                    f"{escape_markdown(unit.category)} |"
                )
                for unit in record.units
            )
            word_values = [unit.words for unit in record.units]
            lines.extend(
                [
                    "",
                    f"Word statistics per {label.casefold()}: minimum={min(word_values)}, "
                    f"maximum={max(word_values)}, median={statistics.median(word_values):g}, "
                    f"average={statistics.mean(word_values):.2f}.",
                    "",
                ]
            )
    return lines


def unit_column_label(extension: str) -> str:
    if extension == ".xlsx":
        return "Sheet"
    if extension == ".pptx":
        return "Slide"
    return "Page"


def configuration_fingerprint() -> str:
    encoded = json.dumps(CONFIGURATION, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()[:16]


def model_inventory_fingerprint() -> str:
    root = Path(os.environ.get("DOCLING_ARTIFACTS_PATH", "/opt/docling/models"))
    digest = hashlib.sha256()
    if not root.is_dir():
        return "missing"
    for path in sorted(root.rglob("*")):
        if path.is_file():
            digest.update(path.relative_to(root).as_posix().encode())
            digest.update(str(path.stat().st_size).encode())
    return digest.hexdigest()[:16]


def tesseract_version() -> str:
    completed = subprocess.run(
        ["/usr/bin/tesseract", "--version"],
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout.splitlines()[0]


def escape_markdown(value: Any) -> str:
    text = str(value).replace("\\", "\\\\").replace("|", "\\|")
    for character in ("`", "*", "_", "[", "]"):
        text = text.replace(character, f"\\{character}")
    return re.sub(r"[\r\n]+", "<br>", text)


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as stream:
            stream.write(content)
            temporary_path = Path(stream.name)
        temporary_path.replace(path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    started_at = datetime.now(timezone.utc)  # noqa: UP017 -- analysis image uses Python 3.10.
    started = time.monotonic()
    try:
        discovered = discover_files(args.input)
        selected = select_files(discovered, args.max_files)
        records = [inspect_file(path, args.input) for path in selected]
        convertible = [
            record
            for record in records
            if record.supported and record.status not in {"failure", "skipped"}
        ]
        if convertible:
            converter = build_converter()
            for index, record in enumerate(convertible, 1):
                print(
                    f"[{index}/{len(convertible)}] {record.relative_path}",
                    file=sys.stderr,
                    flush=True,
                )
                analyze_record(record, converter)

        report = render_report(
            records=records,
            input_label=(
                f"{args.input.name}/"
                if args.input.is_absolute()
                else f"{args.input.as_posix().rstrip('/')}/"
            ),
            max_files=args.max_files,
            started_at=started_at,
            elapsed_seconds=time.monotonic() - started,
        )
        atomic_write(args.output, report)
        print(f"Wrote {args.output}")
        return 0
    except Exception as error:
        print(f"Corpus analysis failed: {type(error).__name__}: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
